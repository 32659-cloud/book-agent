import os
from datetime import datetime

import pytest
from openpyxl import load_workbook

from book_agent.storage import (
    COL,
    BookStore,
    StorageBusyError,
    clean_text,
    safe_filename_part,
)


@pytest.fixture
def store(tmp_path):
    return BookStore(tmp_path / "data")


def add(store, uid=1, title="海邊的卡夫卡", author="村上春樹", category="文學小說", bought=True):
    return store.add_book(uid, f"user{uid}", title, author, category, bought, "photos/x.jpg")


# ---- 防護 3：檔名清洗 ----
@pytest.mark.parametrize(
    "raw",
    ["../../etc/passwd", "a/b\\c", "C:\\Windows\\x", 'a:b*c?d"e<f>g|h', "..", "...", "CON", "  . "],
)
def test_filename_has_no_path_characters(raw):
    s = safe_filename_part(raw)
    for bad in '/\\:*?"<>|':
        assert bad not in s
    assert ".." not in s
    assert s and s == s.strip(" .")
    assert s.upper() != "CON"


def test_filename_empty_becomes_default_and_is_truncated():
    assert safe_filename_part("") == "未命名"
    assert len(safe_filename_part("書" * 100)) == 20


def test_photo_save_stays_inside_user_folder_and_handles_same_second(store):
    now = datetime(2026, 10, 25, 12, 0, 0)
    p1 = store.save_photo(123, "../../evil", ".jpg", b"a", now)
    p2 = store.save_photo(123, "../../evil", ".jpg", b"b", now)
    assert p1 != p2 and p2.endswith("-2.jpg")
    for p in (p1, p2):
        full = (store.data_dir / p).resolve()
        assert (store.photo_dir / "123").resolve() in full.parents
    assert p1.startswith("photos/123/2026-10-25 12-00-00 ")


def test_photo_rejects_non_numeric_user_id(store):
    with pytest.raises(ValueError):
        store.save_photo("../x", "t", ".jpg", b"a")


# ---- 防護 4：Excel 公式注入 ----
@pytest.mark.parametrize("evil", ["=1+1", "+1", "-1", "@SUM(A1)", '=HYPERLINK("http://x","y")'])
def test_formula_text_is_stored_as_plain_text(store, evil):
    add(store, title=evil, author=evil)
    ws = load_workbook(store.xlsx)["書單"]
    for name in ("書名", "作者"):
        cell = ws.cell(row=2, column=COL[name])
        assert cell.data_type == "s"
        assert str(cell.value).startswith("'")
    # 讀回來時是原本的文字
    assert store.list_books(1)[0].title == evil


def test_control_characters_do_not_crash(store):
    add(store, title="a\x00b\x07c\nd")
    assert store.list_books(1)[0].title == "a b c d"


# ---- 紀錄 ----
def test_rows_are_appended_not_overwritten(store):
    add(store, title="A")
    add(store, title="B")
    ws = load_workbook(store.xlsx)["書單"]
    assert ws.max_row == 3
    assert [b.number for b in store.list_books(1)] == [1, 2]


def test_numbers_are_per_user(store):
    assert add(store, 1, "A").number == 1
    assert add(store, 2, "B").number == 1
    assert add(store, 1, "C").number == 2


def test_user_id_is_stored_as_text_without_precision_loss(store):
    uid = 123456789012345678
    add(store, uid)
    assert store.list_books(uid)[0].user_id == "123456789012345678"


def test_users_cannot_see_or_change_each_other(store):
    add(store, 1, "A")
    add(store, 2, "B")
    assert [b.title for b in store.list_books(1)] == ["A"]
    assert store.get_book(2, 1).title == "B"
    assert store.update_field(2, 1, "書名", "改").title == "改"
    assert store.get_book(1, 1).title == "A"
    assert store.mark_deleted(1, 5) is None


def test_duplicate_ignores_spaces_case_and_deleted(store):
    add(store, title="Python Crash Course", author="Eric Matthes")
    assert store.find_duplicate(1, "python  crash course", "ERIC MATTHES")
    assert store.find_duplicate(2, "Python Crash Course", "Eric Matthes") is None
    store.mark_deleted(1, 1)
    assert store.find_duplicate(1, "Python Crash Course", "Eric Matthes") is None


def test_delete_hides_but_keeps_row_and_number_is_not_reused(store):
    add(store, title="A")
    store.mark_deleted(1, 1)
    assert store.list_books(1) == []
    assert load_workbook(store.xlsx)["書單"].max_row == 2
    assert add(store, title="A").number == 2
    assert store.stats(1)["total"] == 1


def test_stats_and_unknown_category(store):
    add(store, title="A", category="其他類")
    add(store, title="B", category="文學小說", bought=False)
    store.update_field(1, 1, "閱讀狀態", "已讀")
    s = store.stats(1)
    assert (s["total"], s["read"], s["unread"], s["bought"], s["not_bought"]) == (2, 1, 1, 1, 1)
    assert s["by_category"] == {"其他": 1, "文學小說": 1}


def test_clean_text():
    assert clean_text("  a \t b\n\nc  ", 50) == "a b c"
    assert clean_text(None, 5) == ""
    assert clean_text("x" * 10, 3) == "xxx"


def test_set_page_writes_a_number_and_validates(store):
    add(store)
    assert store.set_page(1, 1, 120).page == 120
    ws = load_workbook(store.xlsx)["書單"]
    cell = ws.cell(row=2, column=COL["目前頁數"])
    assert cell.value == 120 and cell.data_type == "n"
    assert store.set_page(1, 9, 5) is None
    for bad in (-1, 100000, "5", True):
        with pytest.raises(ValueError):
            store.set_page(1, 1, bad)


def test_old_workbook_without_page_column_still_works(tmp_path):
    from openpyxl import Workbook

    from book_agent.storage import HEADERS

    data = tmp_path / "data"
    data.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "書單"
    ws.append(HEADERS[:-1])  # 舊格式：沒有「目前頁數」欄
    ws.append([1, "2026-10-25 10:00:00", "7", "舊使用者", "舊書", "舊作者", "其他", "未讀", "已購買", "", "photos/7/x.jpg"])
    wb.save(data / "books.xlsx")

    s = BookStore(data)
    assert s.list_books(7)[0].page == 0
    assert s.set_page(7, 1, 33).page == 33
    assert add(s, 7, "新書").number == 2
    assert load_workbook(s.xlsx)["書單"].cell(row=1, column=COL["目前頁數"]).value == "目前頁數"


def test_add_book_with_page(store):
    b = store.add_book(1, "u", "A", "a", "其他", True, "p", page=84)
    assert b.page == 84
    assert store.add_book(1, "u", "B", "b", "其他", True, "p", page=-5).page == 0
    assert store.add_book(1, "u", "C", "c", "其他", True, "p", page=10**9).page == 0


# ---- Excel 被打開時（Windows 會鎖檔，無法換掉正式檔）----
@pytest.mark.skipif(os.name != "nt", reason="只有 Windows 會鎖住正被打開的檔案")
def test_busy_excel_warns_retries_and_raises_without_crashing(store, capsys):
    add(store, title="A")  # 先讓 books.xlsx 存在
    with open(store.xlsx, "rb"):  # 模擬 Excel 正開著這個檔案
        wb, _ = store._open()
        with pytest.raises(StorageBusyError):
            store._save(wb, retries=3, wait=0.01)
    out = capsys.readouterr().out
    assert out.count("[警告]") == 3 and "books.xlsx" in out and "重試" in out
    assert not (store.xlsx.parent / "books.xlsx.tmp").exists()  # 不留下暫存檔
    assert [b.title for b in store.list_books(1)] == ["A"]  # 原本的紀錄沒壞


@pytest.mark.skipif(os.name != "nt", reason="只有 Windows 會鎖住正被打開的檔案")
def test_excel_closed_during_retry_then_write_succeeds(store, capsys):
    import threading

    add(store, title="A")
    handle = open(store.xlsx, "rb")
    threading.Timer(0.25, handle.close).start()  # 0.25 秒後「關掉 Excel」
    b = store.add_book(1, "u", "B", "b", "其他", True, "p")  # 第一次寫入失敗，等 2 秒後重試成功
    assert b.number == 2
    assert [x.title for x in store.list_books(1)] == ["A", "B"]
    assert "[警告]" in capsys.readouterr().out
